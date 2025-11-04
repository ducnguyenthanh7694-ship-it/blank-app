
import io
import re
import base64
from datetime import datetime

import streamlit as st
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from PIL import Image

# =========================
# Helper functions
# =========================

def format_vn(num, decimals=0):
    """
    Format a number using dot as thousand separators.
    If decimals > 0, use comma as decimal separator (optional aesthetic).
    """
    try:
        if decimals > 0:
            s = f"{num:,.{decimals}f}"
            s = s.replace(",", "X").replace(".", ",").replace("X", ".")
            return s
        else:
            s = f"{float(num):,.0f}"
            return s.replace(",", ".")
    except Exception:
        return str(num)

def read_docx_text(file_like):
    try:
        doc = Document(file_like)
        text = []
        for p in doc.paragraphs:
            text.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text.append(cell.text)
        return "\n".join(text)
    except PackageNotFoundError:
        return ""

def naive_parse_from_docx(text):
    """
    Very lightweight regex parsing from uploaded .docx content.
    Tries to extract common fields if they appear in the text.
    """
    result = {}

    # Name (Vietnamese: Họ và tên: ... )
    m = re.search(r"Họ\s*và\s*tên\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    if m:
        result["ho_ten"] = m.group(1).strip()

    # CCCD or CMND
    m = re.search(r"(?:CCCD|CMND)\s*[:\-]?\s*([0-9]{9,12})", text, re.IGNORECASE)
    if m:
        result["cccd"] = m.group(1).strip()

    # Phone
    m = re.search(r"(?:SĐT|Điện thoại|Phone)\s*[:\-]?\s*(\+?[\d\s\-]{8,15})", text, re.IGNORECASE)
    if m:
        result["sdt"] = re.sub(r"\s+", "", m.group(1))

    # Address
    m = re.search(r"(?:Địa chỉ|Address)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    if m:
        result["dia_chi"] = m.group(1).strip()

    # Loan purpose
    m = re.search(r"(?:Mục đích vay|Loan purpose)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    if m:
        result["muc_dich_vay"] = m.group(1).strip()

    # Numeric helpers
    def find_money(label_regex):
        m = re.search(label_regex + r"\s*[:\-]?\s*([\d\. ,]+)", text, re.IGNORECASE)
        if m:
            return to_number(m.group(1))
        return None

    result.setdefault("tong_nhu_cau_von", find_money(r"(?:Tổng nhu cầu vốn|Total capital need)"))
    result.setdefault("von_doi_ung", find_money(r"(?:Vốn đối ứng|Counterpart fund)"))
    result.setdefault("so_tien_vay", find_money(r"(?:Số tiền vay|Loan amount)"))
    result.setdefault("lai_suat", find_money(r"(?:Lãi suất|Interest rate).*?(?:%/năm|per year)?"))
    result.setdefault("thoi_gian_vay", find_money(r"(?:Thời gian vay|Loan tenor).*?(?:tháng|months)?"))

    # Collateral
    m = re.search(r"(?:Mô tả tài sản|Tài sản đảm bảo|Collateral).*?:\s*(.+)", text, re.IGNORECASE)
    if m:
        result["mo_ta_ts"] = m.group(1).strip()

    result.setdefault("gia_tri_ts", find_money(r"(?:Giá trị định giá|Appraised value)"))

    # Cleanup rates and tenor
    if "lai_suat" in result and result["lai_suat"] is not None and result["lai_suat"] > 100:
        result["lai_suat"] = result["lai_suat"] / 100  # In case it's given like 1200 -> 12%

    if "thoi_gian_vay" in result and result["thoi_gian_vay"] is not None and result["thoi_gian_vay"] < 1:
        result["thoi_gian_vay"] = int(round(result["thoi_gian_vay"] * 12))

    return result

def to_number(s):
    if s is None:
        return None
    s = str(s).strip()
    # Normalize: remove dot thousand-sep and spaces, convert comma decimal to dot
    s = s.replace(".", "").replace(" ", "").replace("\u00A0", "")
    s = s.replace(",", ".")
    try:
        val = float(re.sub(r"[^\d\.]", "", s))
        return val
    except Exception:
        return None

def parse_from_image_filename(file_name):
    """
    Very simple heuristics to grab info from image filename:
    e.g., 'Nguyen_Thanh_Duc_CCCD_012345678901.jpg' -> name + CCCD
    e.g., 'STK_500000000.jpg' -> maybe savings amount or collateral value.
    """
    info = {}

    # Extract CCCD if 9-12 consecutive digits appear
    m = re.search(r"(\d{9,12})", file_name)
    if m:
        info["cccd"] = m.group(1)

    # Try to parse a name from underscores/dashes (skip tokens like CCCD, CMND, STK)
    tokens = re.split(r"[._\- ]+", file_name)
    name_tokens = [t for t in tokens if t and t.upper() not in {"JPG", "JPEG", "PNG", "CCCD", "CMND", "STK", "SOTIETKIEM", "SOTK"} and not re.fullmatch(r"\d{9,12}", t)]
    if len(name_tokens) >= 2:
        # Capitalize nicely
        info["ho_ten"] = " ".join([t.capitalize() for t in name_tokens[:3]])  # up to 3 tokens for name

    # If filename contains STK and a large number, treat as appraised value placeholder
    if any(tag in [t.upper() for t in tokens] for tag in ["STK", "SOTK", "SOTIETKIEM"]):
        m2 = re.search(r"(\d{7,12})", file_name)
        if m2:
            info["gia_tri_ts"] = float(m2.group(1))

    return info

def replace_placeholders_in_docx(template_file, mapping):
    """
    Replace {{placeholders}} in template .docx (paragraphs and table cells).
    Returns a bytes buffer of the new document.
    """
    doc = Document(template_file)

    def replace_in_run(run):
        for key, val in mapping.items():
            placeholder = "{{" + key + "}}"
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(val))

    # Replace in paragraphs
    for p in doc.paragraphs:
        for run in p.runs:
            replace_in_run(run)

    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        replace_in_run(run)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def download_button_bytesio(buf, filename, label):
    b64 = base64.b64encode(buf.getvalue()).decode()
    href = f'<a href="data:application/octet-stream;base64,{b64}" download="{filename}">{label}</a>'
    st.markdown(href, unsafe_allow_html=True)

# =========================
# Streamlit App
# =========================

st.set_page_config(
    page_title="Hệ thống làm hồ sơ cầm cố sổ tiết kiệm",
    page_icon="💼",
    layout="wide"
)

st.title("💼 Hệ thống làm hồ sơ cầm cố sổ tiết kiệm")

with st.sidebar:
    st.header("Cấu hình & Xuất dữ liệu")
    api_key = st.text_input("OpenAI API Key", type="password", help="Dùng cho Chatbot hỗ trợ (mặc định không bắt buộc).")
    st.divider()

    export_choice = st.selectbox(
        "Chức năng Xuất dữ liệu",
        ["-- Chọn --", "Xuất báo cáo đề xuất kiêm hợp đồng tín dụng", "Xuất phụ lục nhận tiền vay"]
    )
    export_button = st.button("Thực hiện")

# Session state defaults
if "form" not in st.session_state:
    st.session_state.form = {
        "ho_ten": "",
        "cccd": "",
        "dia_chi": "",
        "sdt": "",
        "muc_dich_vay": "",
        "tong_nhu_cau_von": 0.0,
        "von_doi_ung": 0.0,
        "so_tien_vay": 0.0,
        "lai_suat": 12.0,
        "thoi_gian_vay": 12,
        "mo_ta_ts": "Sổ tiết kiệm",
        "gia_tri_ts": 0.0,
    }

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []  # list of dict(role, content)

# Tabs
tab1, tab2, tab3 = st.tabs(["Nhập liệu & Trích xuất thông tin", "Phân tích Chỉ số & Dòng tiền", "Chatbot Hỗ trợ"])

with tab1:
    st.subheader("Tải tệp & Trích xuất thông tin ban đầu")

    col_docx, col_img = st.columns(2)
    with col_docx:
        docx_file = st.file_uploader("Upload mẫu .docx (Báo cáo/Hợp đồng hoặc Phụ lục)", type=["docx"], accept_multiple_files=False, help="Tệp mẫu sẽ được dùng để thay thế placeholder.")
        if docx_file is not None:
            # Attempt to parse any useful info from existing content to pre-fill fields
            text = read_docx_text(docx_file)
            parsed = naive_parse_from_docx(text)
            st.caption("Đã trích xuất sơ bộ từ .docx (nếu có):")
            if parsed:
                st.json(parsed)
                # Update defaults without overwriting already edited values if those are non-empty / non-zero
                for k, v in parsed.items():
                    if k in st.session_state.form:
                        if isinstance(st.session_state.form[k], (int, float)):
                            if st.session_state.form[k] == 0 or st.session_state.form[k] == 0.0:
                                st.session_state.form[k] = v
                        else:
                            if not st.session_state.form[k]:
                                st.session_state.form[k] = v
            else:
                st.info("Không phát hiện được trường thông tin rõ ràng trong .docx.")

    with col_img:
        img_file = st.file_uploader("Upload ảnh .jpg (sổ tiết kiệm / CCCD)", type=["jpg", "jpeg"], accept_multiple_files=False)
        if img_file is not None:
            try:
                img = Image.open(img_file)
                st.image(img, caption="Ảnh đã tải lên", use_column_width=True)
            except Exception:
                st.warning("Không thể hiển thị ảnh, nhưng vẫn có thể dùng tên file để suy luận sơ bộ.")

            inferred = parse_from_image_filename(img_file.name)
            st.caption("Suy luận đơn giản từ tên file ảnh:")
            if inferred:
                st.json(inferred)
                for k, v in inferred.items():
                    if k in st.session_state.form and (not st.session_state.form[k] or st.session_state.form[k] == 0):
                        st.session_state.form[k] = v
            else:
                st.info("Chưa suy luận được dữ liệu từ tên file. Bạn có thể nhập thủ công.")

    st.subheader("Nhập liệu chi tiết (có thể chỉnh sửa)")

    with st.form("form_inputs", clear_on_submit=False):
        # Vùng 1 - Thông tin khách hàng
        st.markdown("### Vùng 1 - Thông tin khách hàng")
        c1, c2 = st.columns(2)
        with c1:
            ho_ten = st.text_input("Họ và tên", value=st.session_state.form["ho_ten"])
            cccd = st.text_input("CCCD", value=st.session_state.form["cccd"])
        with c2:
            dia_chi = st.text_input("Địa chỉ", value=st.session_state.form["dia_chi"])
            sdt = st.text_input("SĐT", value=st.session_state.form["sdt"])

        # Vùng 2 - Thông tin phương án vay
        st.markdown("### Vùng 2 - Thông tin phương án vay")
        c3, c4, c5 = st.columns(3)
        with c3:
            muc_dich_vay = st.text_input("Mục đích vay", value=st.session_state.form["muc_dich_vay"])
            tong_nhu_cau_von = st.number_input("Tổng nhu cầu vốn (VND)", min_value=0.0, step=1_000_000.0, value=float(st.session_state.form["tong_nhu_cau_von"]), format="%.0f")
        with c4:
            von_doi_ung = st.number_input("Vốn đối ứng (VND)", min_value=0.0, step=1_000_000.0, value=float(st.session_state.form["von_doi_ung"]), format="%.0f")
            so_tien_vay = st.number_input("Số tiền vay (VND)", min_value=0.0, step=1_000_000.0, value=float(st.session_state.form["so_tien_vay"]), format="%.0f")
        with c5:
            lai_suat = st.number_input("Lãi suất (%/năm)", min_value=0.0, max_value=100.0, step=0.1, value=float(st.session_state.form["lai_suat"]), format="%.2f")
            thoi_gian_vay = st.number_input("Thời gian vay (tháng)", min_value=1, step=1, value=int(st.session_state.form["thoi_gian_vay"]), format="%d")

        # Vùng 3 - Thông tin tài sản đảm bảo
        st.markdown("### Vùng 3 - Thông tin tài sản đảm bảo")
        c6, c7 = st.columns([2,1])
        with c6:
            mo_ta_ts = st.text_input("Mô tả tài sản", value=st.session_state.form["mo_ta_ts"])
        with c7:
            gia_tri_ts = st.number_input("Giá trị định giá (VND)", min_value=0.0, step=1_000_000.0, value=float(st.session_state.form["gia_tri_ts"]), format="%.0f")

        submitted = st.form_submit_button("Lưu/Áp dụng dữ liệu")
        if submitted:
            st.session_state.form.update({
                "ho_ten": ho_ten,
                "cccd": cccd,
                "dia_chi": dia_chi,
                "sdt": sdt,
                "muc_dich_vay": muc_dich_vay,
                "tong_nhu_cau_von": float(tong_nhu_cau_von),
                "von_doi_ung": float(von_doi_ung),
                "so_tien_vay": float(so_tien_vay),
                "lai_suat": float(lai_suat),
                "thoi_gian_vay": int(thoi_gian_vay),
                "mo_ta_ts": mo_ta_ts,
                "gia_tri_ts": float(gia_tri_ts),
            })
            st.success("Đã cập nhật dữ liệu. Các tab khác sẽ tự động tính toán lại.")

with tab2:
    st.subheader("Phân tích Chỉ số & Dòng tiền")

    f = st.session_state.form
    # Key ratios
    tong_nhu_cau = f.get("tong_nhu_cau_von", 0.0) or 0.0
    vay = f.get("so_tien_vay", 0.0) or 0.0
    doi_ung = f.get("von_doi_ung", 0.0) or 0.0

    ty_le_vay_tren_nhu_cau = (vay / tong_nhu_cau * 100) if tong_nhu_cau > 0 else 0.0
    ty_le_von_doi_ung = (doi_ung / tong_nhu_cau * 100) if tong_nhu_cau > 0 else 0.0

    colA, colB, colC = st.columns(3)
    colA.metric("Tổng nhu cầu vốn", f"{format_vn(tong_nhu_cau)} VND")
    colB.metric("Số tiền vay", f"{format_vn(vay)} VND")
    colC.metric("Vốn đối ứng", f"{format_vn(doi_ung)} VND")

    colD, colE = st.columns(2)
    colD.metric("Tỷ lệ Vay/Tổng nhu cầu vốn", f"{ty_le_vay_tren_nhu_cau:.2f}%")
    colE.metric("Tỷ lệ Vốn đối ứng", f"{ty_le_von_doi_ung:.2f}%")

    # Optional simple cashflow: monthly interest-only (illustrative)
    st.markdown("#### Minh họa dòng tiền (trả lãi hàng tháng, gốc cuối kỳ)")
    lai_suat_nam = f.get("lai_suat", 0.0) or 0.0
    thang = int(f.get("thoi_gian_vay", 0) or 0)
    if vay > 0 and lai_suat_nam > 0 and thang > 0:
        lai_thang = vay * (lai_suat_nam/100) / 12
        st.write(f"- Lãi hàng tháng ước tính: **{format_vn(lai_thang)} VND**")
        st.write(f"- Tổng lãi ước tính (không tính lãi trên lãi): **{format_vn(lai_thang*thang)} VND**")
        st.write(f"- Gốc thanh toán cuối kỳ: **{format_vn(vay)} VND**")
    else:
        st.info("Nhập đủ Số tiền vay, Lãi suất và Thời gian vay để xem minh họa dòng tiền.")

with tab3:
    st.subheader("Chatbot Hỗ trợ")

    st.caption("Mẹo: Cung cấp bối cảnh như hồ sơ khách hàng, phương án vay để nhận tư vấn.")
    reset = st.button("🗑️ Xóa lịch sử trò chuyện")
    if reset:
        st.session_state.chat_history = []
        st.success("Đã xóa lịch sử.")

    # Show conversation
    for msg in st.session_state.chat_history:
        if msg["role"] == "user":
            st.chat_message("user").write(msg["content"])
        else:
            st.chat_message("assistant").write(msg["content"])

    # Chat input
    user_prompt = st.chat_input("Nhập câu hỏi cho ChatGPT...")
    if user_prompt:
        st.session_state.chat_history.append({"role": "user", "content": user_prompt})
        response_text = "Vui lòng nhập OpenAI API Key ở thanh bên để sử dụng Chatbot."
        if api_key:
            try:
                # Lazy import to avoid hard dependency if not used
                from openai import OpenAI
                client = OpenAI(api_key=api_key)
                completion = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": "Bạn là trợ lý tín dụng ngân hàng, am hiểu thẩm định cho vay cầm cố sổ tiết kiệm. Trả lời súc tích, đúng chuẩn mực ngân hàng."},
                        *st.session_state.chat_history
                    ],
                    temperature=0.2,
                )
                response_text = completion.choices[0].message.content
            except Exception as e:
                response_text = f"Lỗi gọi OpenAI API: {e}"
        st.session_state.chat_history.append({"role": "assistant", "content": response_text})
        st.chat_message("assistant").write(response_text)

# =========================
# Export handling (Sidebar)
# =========================

if export_button:
    f = st.session_state.form.copy()

    # Build mapping for placeholders in templates
    mapping = {
        "ho_ten": f.get("ho_ten", ""),
        "cccd": f.get("cccd", ""),
        "dia_chi": f.get("dia_chi", ""),
        "sdt": f.get("sdt", ""),
        "muc_dich_vay": f.get("muc_dich_vay", ""),
        "tong_nhu_cau_von": format_vn(f.get("tong_nhu_cau_von", 0.0)),
        "von_doi_ung": format_vn(f.get("von_doi_ung", 0.0)),
        "so_tien_vay": format_vn(f.get("so_tien_vay", 0.0)),
        "lai_suat": f"{f.get('lai_suat', 0):.2f}%/năm",
        "thoi_gian_vay": str(f.get("thoi_gian_vay", "")),
        "mo_ta_ts": f.get("mo_ta_ts", ""),
        "gia_tri_ts": format_vn(f.get("gia_tri_ts", 0.0)),
        # Derived
        "ty_le_vay_tren_nhu_cau": f"{(f.get('so_tien_vay',0)/f.get('tong_nhu_cau_von',1)*100 if f.get('tong_nhu_cau_von',0)>0 else 0):.2f}%",
        "ty_le_von_doi_ung": f"{(f.get('von_doi_ung',0)/f.get('tong_nhu_cau_von',1)*100 if f.get('tong_nhu_cau_von',0)>0 else 0):.2f}%",
        "ngay_lap": datetime.now().strftime("%d/%m/%Y"),
    }

    if export_choice == "-- Chọn --":
        st.warning("Vui lòng chọn loại tài liệu cần xuất.")
    else:
        # Need a template .docx to export
        if 'docx_file' in locals() and docx_file is not None:
            try:
                outbuf = replace_placeholders_in_docx(docx_file, mapping)
                if "Báo cáo" in export_choice or "hợp đồng" in export_choice.lower():
                    filename = "bao_cao_de_xuat_hop_dong_da_dien.docx"
                else:
                    filename = "phu_luc_nhan_tien_vay_da_dien.docx"
                st.success("Đã tạo file. Nhấn link bên dưới để tải về:")
                download_button_bytesio(outbuf, filename, f"⬇️ Tải {filename}")
            except Exception as e:
                st.error(f"Lỗi khi xử lý template .docx: {e}")
        else:
            st.error("Chưa có file mẫu .docx. Hãy upload ở tab 'Nhập liệu & Trích xuất thông tin'.")
